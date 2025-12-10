from pptx import Presentation
from pptx.dml.color import RGBColor

from Builders.charts import add_chart_element
from Builders.shapes import add_bg_shape, add_image_element, add_pseudo_element, add_shape_element
from Builders.tables import add_table_element
from Builders.texts import add_text_element, add_list_element, add_inline_group_element, add_link_element
from Builders.extractors import extract_table_info_from_hierarchy, extract_list_info_from_hierarchy, extract_inline_group_from_hierarchy

from Builders.utils import pixels_to_emu, safe_int, parse_color, has_any_border
from Builders.logger import logger

def _get_slides_functions():
    from Builders.slides import parse_images, prefetch_images, add_slide, create_pptx_from_json
    return parse_images, prefetch_images, add_slide, create_pptx_from_json


def create_pptx_from_slides_data(slides_data, reference_id=0):
    """
    Create a pptx Presentation object from slide data directly.

    Prefetches images, sets presentation dimensions based on the first slide,
    and iterates slides to build the presentation. This function handles
    hierarchical slide data structure.

    Parameters:
        slides_data (list): List of slide data dictionaries.
        reference_id (str): Identifier used for logging and image fetching.

    Returns:
        pptx.presentation.Presentation: The constructed Presentation object or None on failure.
    """
    if not slides_data:
        logger.warning("No slides found in data", reference_id=reference_id)
        return None
    
    parse_images, prefetch_images, add_slide, _ = _get_slides_functions()
    
    # Prefetch images for better performance
    images_to_fetch = parse_images(slides_data)
    img_dict = prefetch_images(images_to_fetch, reference_id)

    first_slide = slides_data[0]
    slide_width = safe_int(first_slide.get('slideWidth', 1920))
    slide_height = safe_int(first_slide.get('slideHeight', 1080))

    # Create presentation with precise dimensions
    prs = Presentation()
    prs.slide_width = pixels_to_emu(slide_width)
    prs.slide_height = pixels_to_emu(slide_height)
    
    for s_idx, slide_data in enumerate(slides_data):
        try:
            # Use add_slide for flat element structure, otherwise process hierarchically
            if 'elements' in slide_data:
                add_slide(prs, slide_data, slide_width, slide_height, reference_id, img_dict)
            else:
                # Hierarchical structure with 'children'
                slide_layout = prs.slide_layouts[6]  # Blank layout
                slide = prs.slides.add_slide(slide_layout)
                
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = parse_color('#ffffff') or RGBColor(255, 255, 255)
                
                slide_styles = slide_data.get('slideStyles', {})
                if slide_styles:
                    add_bg_shape(slide, slide_styles, 0, 0, slide_width, slide_height, reference_id=reference_id)
                
                children = slide_data.get('children', [])
                elements_sorted = sorted(children, key=lambda e: (e.get('zIndex', 0), e.get('y', 0), e.get('x', 0)))
                
                processed_element_ids = set()
                processed_text_content = set()
                table_positions = {}
                
                for element in elements_sorted:
                    element_id = f"{element.get('type', 'unknown')}-{element.get('x', 0)}-{element.get('y', 0)}-{element.get('className', '')}"
                    if element_id in processed_element_ids:
                        continue
                    processed_element_ids.add(element_id)
                    
                    try:
                        process_element_recursive(slide, element, slide_width, slide_height, processed_text_content, table_positions, img_dict=img_dict, reference_id=reference_id)
                    except Exception as e:
                        logger.error(f"Error processing element in slide {s_idx}: {e}", reference_id=reference_id)
        except Exception as e:
            logger.error(f"There was an exception in slide idx: {s_idx} due to {e}", reference_id=reference_id)

    return prs

def process_element_recursive(slide, element, slide_width, slide_height, processed_text_content, table_positions, parent_has_shadow=False, img_dict={}, reference_id=0):
    """Recursively process element hierarchy from Extract.js structure.
    
    Parameters:
        slide: pptx Slide object.
        element (dict): Element to process.
        slide_width (int): Slide width in pixels.
        slide_height (int): Slide height in pixels.
        processed_text_content (set): Set of already processed text content identifiers.
        table_positions (dict): Dict storing table positions for overlay alignment.
        parent_has_shadow (bool): Whether parent has shadow styling.
        img_dict (dict): Cache of prefetched images.
        reference_id (int|str): Identifier for logging/tracking.
    """
    element_type = element.get('type')
    element_class = element.get('className', '')
    
    # Process pseudo-before if present
    if element.get('pseudoBefore'):
        pseudo_before = element['pseudoBefore']
        pseudo_before['type'] = 'pseudo'
        pseudo_before['pseudoType'] = '::before'
        add_pseudo_element(slide, pseudo_before, slide_width, slide_height, reference_id=reference_id)
    
    # Handle complex shapes first (includes arrows) - prevent text duplication
    if element.get('shapeInfo'):
        shape_text = element.get('shapeInfo', {}).get('text', element.get('text', '')).strip()
        shape_class = element.get('className', '')
        
        # Create a unique identifier for this text content + class + position
        text_class_id = f"{shape_text}-{shape_class}-{element.get('x', 0)}-{element.get('y', 0)}"
        
        # Skip if this exact text with same class at similar position was already processed
        if shape_text and text_class_id in processed_text_content:
            return
        
        # Generic container/child duplicate detection
        if 'container' in shape_class:
            # Check children for same text
            has_child_element = False
            for child in element.get('children', []):
                if child.get('shapeInfo', {}).get('text', '').strip() == shape_text and \
                   'container' not in child.get('className', ''):
                    has_child_element = True
                    break
            
            if has_child_element:
                # Process children instead
                for child in element.get('children', []):
                    process_element_recursive(slide, child, slide_width, slide_height, processed_text_content, table_positions, parent_has_shadow, img_dict=img_dict, reference_id=reference_id)
                return
        
        # Add to processed text content
        if shape_text:
            processed_text_content.add(text_class_id)
        
        add_shape_element(slide, element)
        # Don't process children for shapes
        return
    
    # Process element based on type
    if element_type == 'table':
        # Extract table info from hierarchical structure
        table_info = extract_table_info_from_hierarchy(element)
        element['tableInfo'] = table_info
        
        # Store table position
        table_key = f"table_{element.get('x', 0)}_{element.get('y', 0)}"
        table_positions[table_key] = {
            'x': element.get('x', 0),
            'y': element.get('y', 0),
            'table_info': table_info
        }
        
        add_table_element(slide, element, slide_width, slide_height, reference_id, parent_has_shadow, img_dict=img_dict)
        # Table processing handles its own children
        return
    
    elif element_type == 'img':
        # Extract media info
        if not element.get('mediaInfo'):
            element['mediaInfo'] = {
                'src': element.get('src', ''),
                'naturalWidth': element.get('naturalWidth', element.get('width', 0)),
                'naturalHeight': element.get('naturalHeight', element.get('height', 0))
            }
        add_image_element(slide, element, slide_width, slide_height, parent_has_shadow, img_dict=img_dict, reference_id=reference_id)
        return
    
    elif element_type == 'a':
        # Extract link info
        if not element.get('linkInfo'):
            element['linkInfo'] = {
                'href': element.get('href', ''),
                'text': element.get('text', '')
            }
        add_link_element(slide, element, slide_width, slide_height, parent_has_shadow)
        return
    
    elif element_type == 'canvas' and element.get('chartInfo'):
        add_chart_element(slide, element, slide_width, slide_height, reference_id=reference_id)
        return
    
    elif element_type in ['ul', 'ol']:
        # Extract list info from hierarchy
        list_info = extract_list_info_from_hierarchy(element)
        element['listInfo'] = list_info
        add_list_element(slide, element, parent_has_shadow=parent_has_shadow, reference_id=reference_id)
        # List processing handles its own children
        return
    
    elif element_type == 'pseudo':
        add_pseudo_element(slide, element, slide_width, slide_height, reference_id=reference_id)
        return
    
    # Check if element has inline children (span, strong, em, etc.)
    children = element.get('children', [])
    inline_types = {'span', 'strong', 'em', 'b', 'i', 'u', 'br', 'a'}
    
    has_inline_children = any(child.get('type') in inline_types for child in children)
    has_non_inline_children = any(child.get('type') not in inline_types for child in children)
    
    if has_inline_children and not has_non_inline_children:
        # Extract inline group
        inline_group = extract_inline_group_from_hierarchy(element)
        element['inlineGroup'] = inline_group
        add_inline_group_element(slide, element, slide_width, slide_height, parent_has_shadow, reference_id=reference_id)
        # Don't process children separately
        return
    
    # For block elements with text
    if element_type in ['div', 'span', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p']:
        if element.get('text', '').strip():
            add_text_element(slide, element, slide_width, slide_height, parent_has_shadow)
        elif has_any_border(element.get('styles', {})) or \
             parse_color(element.get('styles', {}).get('backgroundColor')) or \
             element.get('styles', {}).get('boxShadow', 'none') != 'none':
            x = element.get('x', 0)
            y = element.get('y', 0)
            width = max(1, element.get('width', 100))
            height = max(1, element.get('height', 100))
            add_bg_shape(slide, element.get('styles', {}), x, y, width, height, reference_id=reference_id)
    
    # Process pseudo-after if present
    if element.get('pseudoAfter'):
        pseudo_after = element['pseudoAfter']
        pseudo_after['type'] = 'pseudo'
        pseudo_after['pseudoType'] = '::after'
        add_pseudo_element(slide, pseudo_after, slide_width, slide_height, reference_id=reference_id)
    
    # Recursively process children
    for child in children:
        process_element_recursive(slide, child, slide_width, slide_height, processed_text_content, table_positions, parent_has_shadow, img_dict=img_dict, reference_id=reference_id)


if __name__ == "__main__":
    import json
    _, _, _, create_pptx_from_json = _get_slides_functions()
    
    json_path = 'slides_data.json'
    try:
        with open(json_path, 'r', encoding='utf-8') as f:
            slides_data = json.load(f)
        prs = create_pptx_from_json(slides_data, reference_id=0)
        if prs:
            prs.save('output2.pptx')
            logger.info("Presentation saved as 'output2.pptx'", reference_id=0)
    except Exception as e:
        logger.error(f"Error: {e}", reference_id=0)