from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from Builders.utils import pixels_to_emu, safe_float, safe_int, get_font_size_pt, parse_color, parse_border_radius, is_uniform_border, has_any_border, get_vertical_alignment, get_para_alignment, is_bold, makeParaBulletPointed, px_to_pt
from Builders.shapes import add_bg_shape
from Builders.logger import logger

def add_link_element(slide, element, slide_width, slide_height, parent_has_shadow=False, reference_id=0):
    """Add a hyperlink element to the slide (handles both direct and linkInfo style)"""
    # Support both direct and linkInfo structure
    link_info = element.get('linkInfo', {})
    link_href = link_info.get('href') or element.get('href', '')
    link_text = link_info.get('text') or element.get('text', '').strip()
    if not link_text:
        return

    x = element.get('x', 0)
    y = element.get('y', 0)
    width = max(1, element.get('width', 100))
    height = max(1, element.get('height', 20))
    x = max(0, min(x, slide_width - width))
    y = max(0, min(y, slide_height - height))
    styles = element.get('styles', {})

    textbox = slide.shapes.add_textbox(
        pixels_to_emu(x), pixels_to_emu(y),
        pixels_to_emu(width), pixels_to_emu(height)
    )
    text_frame = textbox.text_frame
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.clear()
    p = text_frame.paragraphs[0]
    run = p.add_run()
    run.text = link_text
    font = run.font
    font_size_px = safe_float(styles.get('fontSize', '12').replace('px', ''))
    font.name = styles.get('fontFamily', 'Arial').split(',')[0].strip('"\'')
    font.size = Pt(max(6, get_font_size_pt(font_size_px)))
    font.bold = styles.get('fontWeight', '400') in ['bold', '600', '700', '800', '900']
    font.italic = styles.get('fontStyle', 'normal') == 'italic'
    link_color = parse_color(styles.get('color', '#0066cc'))
    if link_color:
        font.color.rgb = link_color
    else:
        font.color.rgb = RGBColor(0, 102, 204)
    text_decoration = styles.get('textDecoration', 'underline')
    if 'underline' in text_decoration:
        font.underline = True
    if link_href:
        try:
            hyperlink = run.hyperlink
            hyperlink.address = link_href
        except Exception as hyperlink_error:
            logger.error(f"Failed to add hyperlink functionality: {hyperlink_error}", reference_id=reference_id)
    text_align = styles.get('textAlign', 'left')
    if text_align == 'center':
        p.alignment = PP_ALIGN.CENTER
    elif text_align == 'right':
        p.alignment = PP_ALIGN.RIGHT
    else:
        p.alignment = PP_ALIGN.LEFT
    textbox.fill.background()
    textbox.line.fill.background()
    textbox.shadow.inherit = False

def add_inline_group_element(slide, element, slide_width, slide_height, parent_has_shadow=False, reference_id=0):
    """Render inline group elements (mixed text styles) as textboxes with multiple runs.
    
    Args:
        slide: pptx Slide object.
        element (dict): Element containing inlineGroup info.
        slide_width (int): Width of slide in pixels.
        slide_height (int): Height of slide in pixels.
        parent_has_shadow (bool): Whether parent element has shadow.
        reference_id (int|str): Identifier for diagnostic context.
    """
    inline_group = element.get('inlineGroup')
    if not inline_group:
        return
    inline_elements = inline_group.get('inlineElements', [])
    if not inline_elements:
        return
    
    has_content = any(elem.get('text', '').strip() for elem in inline_elements)
    if not has_content:
        return
    
    group_rect = inline_group.get('groupRect', {})
    x = max(0, min(group_rect.get('x', 0), slide_width - 10))
    y = max(0, min(group_rect.get('y', 0), slide_height - 10))
    width = max(10, min(group_rect.get('width', 100), slide_width - x))
    height = max(10, min(group_rect.get('height', 20), slide_height - y))
    styles = inline_group.get('styles', {})
    
    try:
        # Only add text content if there are inline elements
        if has_content:
            # Add regular background shapes
            if parse_color(styles.get('backgroundColor')) or has_any_border(styles):
                add_bg_shape(slide, styles, x, y, width, height, reference_id=reference_id)
            
            textbox = slide.shapes.add_textbox(
                pixels_to_emu(x), pixels_to_emu(y),
                pixels_to_emu(width), pixels_to_emu(height)
            )
            text_frame = textbox.text_frame
            text_frame.word_wrap = True
            text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            text_frame.margin_left = pixels_to_emu(safe_float(styles.get('paddingLeft', '0px').replace('px', '')))
            text_frame.margin_right = pixels_to_emu(safe_float(styles.get('paddingRight', '0px').replace('px', '')))
            text_frame.margin_top = pixels_to_emu(safe_float(styles.get('paddingTop', '0px').replace('px', '')))
            text_frame.margin_bottom = pixels_to_emu(safe_float(styles.get('paddingBottom', '0px').replace('px', '')))
            textbox.fill.background()
            textbox.line.fill.background()
            textbox.shadow.inherit = False
            text_frame.clear()
            p = text_frame.paragraphs[0]
            first = True
            text_align = styles.get('textAlign', 'left')
            alignment = get_para_alignment(text_align)
            p.alignment = alignment
            
            for inline_element in inline_elements:
                if inline_element.get('type') == 'br':
                    if p is not None:
                        p = text_frame.add_paragraph()
                        p.alignment = alignment
                    first = True
                    continue
                element_text = inline_element.get('text', '')
                if first:
                    element_text = element_text.lstrip()
                if not element_text.strip():
                    continue
                first = False
                run = p.add_run()
                run.text = element_text
                inline_styles = inline_element.get('styles', {})
                font = run.font
                font_size_px = safe_float(inline_styles.get('fontSize', '16').replace('px', ''))
                font.name = inline_styles.get('fontFamily', 'Segoe UI').split(',')[0].strip('"\'')
                font.size = Pt(get_font_size_pt(font_size_px))
                font.bold = is_bold(inline_styles.get('fontWeight'))
                font.italic = inline_styles.get('fontStyle', 'normal') == 'italic'
                color = parse_color(inline_styles.get('color'))
                if color:
                    font.color.rgb = color
            
            # Trim trailing spaces from the last run in the last paragraph
            if text_frame.paragraphs and text_frame.paragraphs[-1].runs:
                last_run = text_frame.paragraphs[-1].runs[-1]
                last_run.text = last_run.text.rstrip()
                
    except Exception as e:
        logger.exception(f"Failed to add inline group element: {e}", reference_id=reference_id)
            

def add_overlay_element(slide, element, slide_width, slide_height, reference_id=0):
    """Render overlay elements as shapes using their CSS.
    
    Args:
        slide: pptx Slide object.
        element (dict): Element containing overlay info.
        slide_width (int): Width of slide in pixels.
        slide_height (int): Height of slide in pixels.
        reference_id (int|str): Identifier for diagnostic context.
    """
    x = element.get('x', 0)
    y = element.get('y', 0)
    width = max(1, element.get('width', 100))
    height = max(1, element.get('height', 10))
    styles = element.get('styles', {})
    
    # For overlay elements, use the rect coordinates if available
    rect = element.get('rect', {})
    if rect:
        x = rect.get('x', x)
        y = rect.get('y', y)
        width = max(1, rect.get('width', width))
        height = max(1, rect.get('height', height))
    
    x = max(0, min(x, slide_width - width))
    y = max(0, min(y, slide_height - height))
    
    bg_color = parse_color(styles.get('backgroundColor'))
    border_radius_str = styles.get('borderRadius', '0px')
    border_radius = parse_border_radius(border_radius_str, width, height)
    has_radius = border_radius > 0
    has_border = is_uniform_border(styles)
    has_any_border_sides = has_any_border(styles)
    box_shadow = styles.get('boxShadow', 'none')
    has_shadow = box_shadow != 'none'
    
    if bg_color or has_border or has_any_border_sides or has_radius or has_shadow:
        add_bg_shape(slide, styles, x, y, width, height, reference_id=reference_id)
    
    text = element.get('text', '').strip()
    if text:
        textbox = slide.shapes.add_textbox(
            pixels_to_emu(x), pixels_to_emu(y),
            pixels_to_emu(width), pixels_to_emu(height)
        )
        text_frame = textbox.text_frame
        text_frame.word_wrap = True
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        text_frame.clear()
        p = text_frame.paragraphs[0]
        run = p.add_run()
        run.text = text
        font = run.font
        font_size_px = safe_float(styles.get('fontSize', '12'))
        font.name = styles.get('fontFamily', 'Arial').split(',')[0].strip('"\'')
        font.size = Pt(max(6, get_font_size_pt(font_size_px)))
        font.bold = is_bold(styles.get('fontWeight'))
        font.italic = styles.get('fontStyle') == 'italic'
        color = parse_color(styles.get('color', 'black'))
        if color:
            font.color.rgb = color
        textbox.fill.background()
        textbox.line.fill.background()
        textbox.shadow.inherit = False

def add_list_paragraphs(text_frame, list_info, level=0, counters=None, element=None, bullet_style_override=None):
    if counters is None:
        counters = {}
        
    list_type = list_info.get('type')
    is_ordered = list_type == 'ol'

    if is_ordered:
        counter_key = f'ol_{level}'
        counters[counter_key] = list_info.get('start', 1) - 1
    
    items = list_info.get('items', [])
    avg_space_px = 0
    if len(items) > 1:
        spaces = []
        for i in range(len(items) - 1):
            item_bottom = items[i]['rect']['y'] + items[i]['rect']['height']
            next_top = items[i+1]['rect']['y']
            space_px = next_top - item_bottom
            if space_px > 0:
                spaces.append(space_px)
        avg_space_px = sum(spaces) / len(spaces) if spaces else 0
    space_after_pt = px_to_pt(avg_space_px)
    line_height_str = list_info.get('styles', {}).get('lineHeight', 'normal')
    if line_height_str == 'normal':
        line_spacing = 1.15
    else:
        try:
            line_spacing = float(line_height_str)
        except ValueError:
            line_spacing = 1.15

    # Determine bullet style for this list
    if bullet_style_override:
        bullet_char, bullet_color, bullet_font, bullet_size_pt = bullet_style_override
    else:
        bullet_char, bullet_color, bullet_font, bullet_size_pt = get_bullet_style_for_list(element or {}, list_info)

    first_item = True
    for item in items:
        p = None
        item_styles = item.get('styles', {})
        default_font_size_px = safe_float(item_styles.get('fontSize', '16').replace('px', ''))
        default_font_size_pt = get_font_size_pt(default_font_size_px)
        default_font_name = item_styles.get('fontFamily', 'Segoe UI').split(',')[0].strip('"\'')
        default_color = parse_color(item_styles.get('color'))

        if item.get('inlineGroup') and item['inlineGroup'].get('inlineElements'):
            inline_elements = item['inlineGroup']['inlineElements']
            first = True
            for inline_element in inline_elements:
                if inline_element.get('type') == 'br':
                    if p is not None:
                        p = text_frame.add_paragraph()
                        # Apply bullet formatting with proper size
                        makeParaBulletPointed(p, bullet_char, bullet_font, bullet_size_pt)
                        p.level = level
                        p.space_after = Pt(space_after_pt)
                        p.line_spacing = line_spacing
                    first = True
                    continue
                element_text = inline_element.get('text', '')
                if p is None:
                    # Use last paragraph if it's empty (no runs), otherwise add new one
                    last_para = text_frame.paragraphs[-1] if text_frame.paragraphs else None
                    if first_item and level == 0 and last_para and len(last_para.runs) == 0:
                        p = last_para
                    else:
                        p = text_frame.add_paragraph()
                    # Apply bullet formatting with proper size
                    makeParaBulletPointed(p, bullet_char, bullet_font, bullet_size_pt)
                    p.level = level
                    p.space_after = Pt(space_after_pt)
                    p.line_spacing = line_spacing
                
                if first:
                    element_text = element_text.lstrip()
                if not element_text.strip():
                    continue
                first = False
                run = p.add_run()
                run.text = element_text
                inline_styles = inline_element.get('styles', {})
                font = run.font
                font_size_px = safe_float(inline_styles.get('fontSize', '16').replace('px', ''))
                font.name = inline_styles.get('fontFamily', 'Segoe UI').split(',')[0].strip('"\'')
                font.size = Pt(get_font_size_pt(font_size_px))
                font.bold = inline_styles.get('fontWeight', '400') in ['bold', '600', '700', '800', '900']
                font.italic = inline_styles.get('fontStyle', 'normal') == 'italic'
                color = parse_color(inline_styles.get('color'))
                if color:
                    font.color.rgb = color
            if text_frame.paragraphs and text_frame.paragraphs[-1].runs:
                last_run = text_frame.paragraphs[-1].runs[-1]
                last_run.text = last_run.text.rstrip()
        else:
            # Use last paragraph if it's empty (no runs), otherwise add new one
            last_para = text_frame.paragraphs[-1] if text_frame.paragraphs else None
            if first_item and level == 0 and last_para and len(last_para.runs) == 0:
                p = last_para
            else:
                p = text_frame.add_paragraph()
            # Apply bullet formatting with proper size
            makeParaBulletPointed(p, bullet_char, bullet_font, bullet_size_pt)
            p.level = level
            p.space_after = Pt(space_after_pt)
            p.line_spacing = line_spacing

            run = p.add_run()
            run.text = item.get('text', '').strip()
            font = run.font
            font.name = default_font_name
            font.size = Pt(default_font_size_pt)
            font.bold = item_styles.get('fontWeight', '400') in ['bold', '600', '700', '800', '900']
            font.italic = item_styles.get('fontStyle', 'normal') == 'italic'
            color = parse_color(item_styles.get('color'))
            if color:
                font.color.rgb = color

        first_item = False
        if item.get('nestedList'):
            add_list_paragraphs(text_frame, item['nestedList'], level + 1, counters, element, bullet_style_override)

def get_bullet_style_for_list(element, list_info):
    """
    Determine bullet char, color, font, and size for a list based on className or CSS styles.
    Returns (bullet_char, bullet_color, bullet_font, bullet_size_pt)
    """
    class_name = element.get('className', '')
    styles = element.get('styles', {})
    bullet_char = '\u2022'
    bullet_color = RGBColor(40, 167, 69)
    bullet_font = 'Arial'
    bullet_size_pt = 12

    # Try to infer from class or parent context
    if 'top-points' in class_name:
        bullet_color = RGBColor(40, 167, 69)
        bullet_font = 'Arial'
        bullet_size_pt = 14
        bullet_char = '●'
    elif 'category-points' in class_name:
        bullet_color = RGBColor(40, 167, 69)
        bullet_font = 'Arial'
        bullet_size_pt = 12
        bullet_char = '●'
    # Try to get from styles if present
    color_str = styles.get('color')
    if color_str:
        parsed = parse_color(color_str)
        if parsed:
            bullet_color = parsed
    font_size = safe_float(styles.get('fontSize', '12'))
    if font_size:
        bullet_size_pt = get_font_size_pt(font_size)
    font_family = styles.get('fontFamily')
    if font_family:
        bullet_font = font_family.split(',')[0].strip('"\'')
    # Try to get bullet char from ::before content if available
    if 'listInfo' in element and element['listInfo'].get('bulletChar'):
        bullet_char = element['listInfo']['bulletChar']
    return bullet_char, bullet_color, bullet_font, bullet_size_pt

def add_list_element(slide, element, slide_width=None, slide_height=None, parent_has_shadow=False, reference_id=0):
    """Add a list block to a slide based on extracted listInfo.

    Creates background shapes if required, a textbox sized to the list rect,
    and calls add_list_paragraphs to populate paragraphs and nested lists.
    
    Args:
        slide: pptx Slide object.
        element (dict): Element containing listInfo.
        slide_width (int): Width of slide in pixels (optional).
        slide_height (int): Height of slide in pixels (optional).
        parent_has_shadow (bool): Whether parent has shadow.
        reference_id (int|str): Identifier for diagnostic context.
    """
    list_info = element.get('listInfo', {})
    if not list_info.get('items'):
        return
    rect = list_info.get('rect', {})
    x = safe_int(rect.get('x', 0))
    y = safe_int(rect.get('y', 0))
    width = safe_int(rect.get('width', element.get('width', 100)))
    height = safe_int(rect.get('height', element.get('height', 100)))
    styles = element.get('styles', {})
    box_shadow = styles.get('boxShadow', 'none')
    has_shadow = box_shadow != 'none' and not parent_has_shadow
    bg_color = parse_color(styles.get('backgroundColor'))
    border_radius_str = styles.get('borderRadius', '0px')
    border_radius = parse_border_radius(border_radius_str, width, height)
    has_radius = border_radius > 0
    has_border = is_uniform_border(styles)
    has_any_border_sides = has_any_border(styles)
    try:
        if bg_color or has_border or has_any_border_sides or has_radius or has_shadow:
            add_bg_shape(slide, styles, x, y, width, height)
        textbox = slide.shapes.add_textbox(
            pixels_to_emu(x), pixels_to_emu(y),
            pixels_to_emu(width), pixels_to_emu(height)
        )
        text_frame = textbox.text_frame
        text_frame.word_wrap = True
        text_frame.vertical_anchor = MSO_ANCHOR.TOP
        text_frame.margin_left = pixels_to_emu(safe_float(styles.get('paddingLeft', '0px').replace('px', '')))
        text_frame.margin_right = pixels_to_emu(safe_float(styles.get('paddingRight', '0px').replace('px', '')))
        text_frame.margin_top = pixels_to_emu(safe_float(styles.get('paddingTop', '0px').replace('px', '')))
        text_frame.margin_bottom = pixels_to_emu(safe_float(styles.get('paddingBottom', '0px').replace('px', '')))
        textbox.fill.background()
        textbox.line.fill.background()
        textbox.shadow.inherit = False
        text_frame.clear()
        bullet_style = get_bullet_style_for_list(element, list_info)
        add_list_paragraphs(text_frame, list_info, element=element, bullet_style_override=bullet_style)
    except Exception as e:
        logger.exception(f"Failed to add list: {e}", reference_id=reference_id)

def add_text_element(slide, element, slide_width, slide_height, parent_has_shadow=False):
    """Enhanced text element creation with precise positioning"""
    text = element.get('text', '').strip()
    if not text:
        return
   
    # Use precise positioning
    x = element.get('x', 0)
    y = element.get('y', 0)
    width = max(1, element.get('width', 100))
    height = max(1, element.get('height', 100))
   
   
    # Ensure coordinates are within slide bounds
    x = max(0, min(x, slide_width - width))
    y = max(0, min(y, slide_height - height))
   
    styles = element.get('styles', {})
    element_class = element.get('className', '')

    if ('table-heading' in element_class or 'heading' in element_class or 
        'pie-heading' in element_class):
        padding_left = safe_float(styles.get('paddingLeft', '0px'))
        
        # For headings with ::before pseudo elements, adjust position
        if padding_left > 0:
            x += max(0, padding_left - 8)   # Smaller offset to prevent too much gap
            width = max(1, width - padding_left + 8)  # Adjust width accordingly
    
    box_shadow = styles.get('boxShadow', 'none')
    has_shadow = box_shadow != 'none' and not parent_has_shadow
    bg_color = parse_color(styles.get('backgroundColor'))
    border_radius_str = styles.get('borderRadius', '0px')
    border_radius = parse_border_radius(border_radius_str, width, height)
    has_radius = border_radius > 0
    has_border = is_uniform_border(styles)
   
    has_any_border_sides = has_any_border(styles)
   
    # Add background/border shapes first
    if bg_color or has_border or has_any_border_sides or has_radius or has_shadow:
        add_bg_shape(slide, styles, x, y, width, height)
    
    # Create text box with precise positioning
    textbox = slide.shapes.add_textbox(
        pixels_to_emu(x), pixels_to_emu(y),
        pixels_to_emu(width), pixels_to_emu(height)
    )
    
    text_frame = textbox.text_frame
    text_frame.word_wrap = True
    
    # Set vertical alignment based on flex properties or default to middle for short text
    display = styles.get('display', 'block')
    align_items = styles.get('alignItems', 'stretch')
    justify_content = styles.get('justifyContent', 'flex-start')
    
    if display == 'flex' and align_items == 'center':
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    elif height < 50: # For small elements like company names
       
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    else:
        text_frame.vertical_anchor = MSO_ANCHOR.TOP
    
    # margin calculation for headings with pseudo elements
    margin_left = safe_float(styles.get('paddingLeft', '0px'))
    margin_right = safe_float(styles.get('paddingRight', '0px'))
    margin_top = safe_float(styles.get('paddingTop', '0px'))
    margin_bottom = safe_float(styles.get('paddingBottom', '0px'))
    
    # Special adjustment for headings with pseudo elements
    if ('table-heading' in element_class or 'pie-heading' in element_class):
        # Reduce left margin to account for positioning adjustment
        margin_left = max(0, margin_left - 12)
    
    text_frame.margin_left = pixels_to_emu(margin_left)
    text_frame.margin_right = pixels_to_emu(margin_right)
    text_frame.margin_top = pixels_to_emu(margin_top)
    text_frame.margin_bottom = pixels_to_emu(margin_bottom)
    
    # Remove textbox styling
    textbox.fill.background()
    textbox.line.fill.background()
    textbox.shadow.inherit = False
    
    # Clear and set text
    text_frame.clear()
    p = text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    
    # Apply text formatting
    font = run.font
    font_size_px = safe_float(styles.get('fontSize', '12'))
    font.name = styles.get('fontFamily', 'Arial').split(',')[0].strip('"\'')
    font.size = Pt(max(6, get_font_size_pt(font_size_px)))
    font.bold = styles.get('fontWeight', '400') in ['bold', '700', '800', '900']
    font.italic = styles.get('fontStyle') == 'italic'
    
    # Apply text color
    color = parse_color(styles.get('color', 'black'))
    if color:
        font.color.rgb = color
    
    # Apply text alignment based on CSS text-align and flex properties
    text_align = styles.get('textAlign', 'left')
    if text_align == 'center' or (display == 'flex' and justify_content == 'center'):
        p.alignment = PP_ALIGN.CENTER
    elif text_align == 'right' or (display == 'flex' and justify_content == 'flex-end'):
        p.alignment = PP_ALIGN.RIGHT
    else:
        p.alignment = PP_ALIGN.LEFT
 